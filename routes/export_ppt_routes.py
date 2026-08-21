from flask import Blueprint, request, jsonify, send_file
from models import db, UserResponse, PricingSKU, TCOEntry
from routes.admin_sections_routes import evaluate_sections
import re, io, copy
from datetime import datetime
from lxml import etree

export_ppt_bp = Blueprint('export_ppt', __name__)

TEMPLATE_PATH = '/home/ubuntu/ztb-super-app/static/assets/zscaler_template.pptx'

def _make_slug(name):
    slug = (name or 'unknown').strip().lower()
    slug = re.sub(r'[^a-z0-9\s-]', '', slug)
    slug = re.sub(r'\s+', '-', slug)
    slug = re.sub(r'-+', '-', slug).strip('-')
    return slug or 'unknown'

def _find_by_slug(customer_slug):
    rows = (UserResponse.query
            .filter_by(status='completed')
            .order_by(UserResponse.completed_at.desc())
            .all())
    for r in rows:
        if _make_slug(r.customer_name) == customer_slug:
            return r
    return None

def _open_template():
    from pptx import Presentation
    from pptx.oxml.ns import qn
    prs = Presentation(TEMPLATE_PATH)
    sl = prs.slides
    for i in range(len(sl) - 1, -1, -1):
        rId = sl._sldIdLst[i].get(qn('r:id'))
        sl._sldIdLst.remove(sl._sldIdLst[i])
        if rId:
            try:
                prs.part.drop_rel(rId)
            except Exception:
                pass
    layout_map = {}
    name_to_layout = {l.name: l for l in prs.slide_layouts}
    blank_navy = name_to_layout.get('Blank [Navy]', prs.slide_layouts[17])
    layout_map['cover']     = blank_navy
    layout_map['title']     = blank_navy
    layout_map['title_sub'] = blank_navy
    layout_map['blank']     = blank_navy
    return prs, layout_map


# ── shared drawing helpers ────────────────────────────────────────────────────

def _box(slide, text, l, t, w, h,
         sz=16, bold=False, color=None, italic=False, align=None):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    if align is None:
        align = PP_ALIGN.LEFT
    txb = slide.shapes.add_textbox(l, t, w, h)
    tf  = txb.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text            = text
    run.font.name       = 'Century Gothic'
    run.font.size       = Pt(sz)
    run.font.bold       = bold
    run.font.italic     = italic
    run.font.color.rgb  = color or WHITE
    return txb

def _rect(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(1, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    return s

def _slide_header(slide, title, subtitle=None):
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    _box(slide, title,
         Emu(370_819), Emu(137_160), Emu(9_274_665), Emu(353_899),
         sz=23, bold=True, color=WHITE)
    if subtitle:
        _box(slide, subtitle,
             Emu(370_705), Emu(548_640), Emu(9_274_665), Emu(250_853),
             sz=16, bold=False, color=WHITE)

def _footer(slide):
    from pptx.util import Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    _box(slide, 'Act Fast. Stay Secure.',
         Emu(370_706), Emu(6_451_135), Emu(2_250_601), Emu(246_220),
         sz=10, color=WHITE)
    _box(slide, '© 2025 Zscaler, Inc. All rights reserved.',
         Emu(9_573_001), Emu(6_472_697), Emu(2_250_601), Emu(223_138),
         sz=8, color=WHITE, align=PP_ALIGN.RIGHT)

def _fmt(v):
    try:    v = float(v)
    except: return '—'
    if v == 0: return '—'
    if v >= 1_000_000: return f'${v/1_000_000:.1f}M'
    if v >= 1_000:     return f'${int(round(v/1000))}K'
    return f'${int(round(v)):,}'

def _fmt_exact(v):
    try:    v = float(v)
    except: return '—'
    if v == 0: return '—'
    return f'${int(round(v)):,}'


def _append_slides_from_file(prs, source_path, slide_index=None):
    from pptx import Presentation
    from lxml import etree
    import copy

    WHITE_HEX = 'FFFFFF'
    A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'

    src = Presentation(source_path)
    slides_to_copy = [src.slides[slide_index]] if slide_index is not None else list(src.slides)

    for src_slide in slides_to_copy:
        blank = prs.slide_layouts[17]
        new_slide = prs.slides.add_slide(blank)

        sp_tree = new_slide.shapes._spTree
        for el in list(sp_tree):
            sp_tree.remove(el)

        for el in src_slide.shapes._spTree:
            tag = el.tag.split('}')[-1] if '}' in el.tag else el.tag
            if tag in ('pic',):
                continue
            cloned = copy.deepcopy(el)
            RPR_TAG = f'{{{A_NS}}}rPr'
            for rpr in cloned.iter(RPR_TAG):
                for color_tag in (f'{{{A_NS}}}solidFill', f'{{{A_NS}}}gradFill',
                                  f'{{{A_NS}}}noFill', f'{{{A_NS}}}pattFill'):
                    for child in rpr.findall(color_tag):
                        rpr.remove(child)
                solid = etree.Element(f'{{{A_NS}}}solidFill')
                srgb  = etree.SubElement(solid, f'{{{A_NS}}}srgbClr')
                srgb.set('val', WHITE_HEX)
                rpr.insert(0, solid)
            END_RPR_TAG = f'{{{A_NS}}}endParaRPr'
            for rpr in cloned.iter(END_RPR_TAG):
                for color_tag in (f'{{{A_NS}}}solidFill', f'{{{A_NS}}}gradFill',
                                  f'{{{A_NS}}}noFill', f'{{{A_NS}}}pattFill'):
                    for child in rpr.findall(color_tag):
                        rpr.remove(child)
                solid = etree.Element(f'{{{A_NS}}}solidFill')
                srgb  = etree.SubElement(solid, f'{{{A_NS}}}srgbClr')
                srgb.set('val', WHITE_HEX)
                rpr.insert(0, solid)
            sp_tree.append(cloned)


def _inject_diagram_png(slide, diagram_path):
    from pptx.util import Emu
    from pptx.dml.color import RGBColor
    import os

    if not os.path.exists(diagram_path):
        return

    try:
        from PIL import Image as PILImage
        with PILImage.open(diagram_path) as im:
            px_w, px_h = im.size
    except Exception:
        px_w, px_h = 1600, 1200

    available_w = Emu(6_400_000)
    available_h = Emu(5_800_000)
    img_l_start = Emu(5_400_000)

    ratio = min(int(available_w) / px_w, int(available_h) / px_h)
    img_w = int(px_w * ratio)
    img_h = int(px_h * ratio)

    img_l = int(img_l_start) + (int(available_w) - img_w) // 2
    img_t = (int(Emu(6_858_000)) - img_h) // 2

    PAD = int(Emu(91_440))
    card = slide.shapes.add_shape(1,
        img_l - PAD, img_t - PAD,
        img_w + PAD * 2, img_h + PAD * 2)
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    card.line.fill.background()

    slide.shapes.add_picture(diagram_path, img_l, img_t, img_w, img_h)


# ══════════════════════════════════════════════════════════════════════════════
# OPPORTUNITY PACKAGE EXPORT
# ══════════════════════════════════════════════════════════════════════════════

@export_ppt_bp.route('/user/<customer_slug>/export-pptx', methods=['POST'])
def export_pptx(customer_slug):
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import matplotlib.ticker as mticker

    user_resp = _find_by_slug(customer_slug)
    if not user_resp:
        return jsonify({'error': 'Not found'}), 404

    data         = request.get_json() or {}
    pricing_rows = data.get('pricing_rows', [])
    phase        = data.get('phase', 'standard')
    support_pct  = float(data.get('support_pct', 20) or 20)
    margin_pct   = float(data.get('margin_pct',  20) or 20)
    services_amt = float(data.get('services_amt', 0) or 0)
    tco_rows     = data.get('tco_rows', [])
    fte_count    = float(data.get('fte_count',   0) or 0)
    fte_cost     = float(data.get('fte_cost',    0) or 0)
    breach_cost  = float(data.get('breach_cost', 0) or 0)
    acv          = float(data.get('acv', 0) or 0)

    section_outcomes = evaluate_sections(user_resp.answers or {})

    NAVY  = RGBColor(0x00, 0x17, 0x44)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BLUE  = RGBColor(0x24, 0x6C, 0xF7)
    CYAN  = RGBColor(0x12, 0xD3, 0xFF)
    RED   = RGBColor(0xED, 0x19, 0x51)
    GREEN = RGBColor(0x6B, 0xFF, 0xB3)
    AMBER = RGBColor(0xFF, 0x93, 0x00)
    MUTED = RGBColor(0xA0, 0xBC, 0xD8)
    CARD  = RGBColor(0x00, 0x23, 0x6B)
    CARD2 = RGBColor(0x00, 0x1F, 0x5E)

    skus      = PricingSKU.query.filter_by(active=True).all()
    app_map   = {s.id: s for s in skus if s.category == 'appliance'}
    sdwan_map = {s.id: s for s in skus if s.category == 'sdwan'}
    seg_map   = {s.id: s for s in skus if s.category == 'segmentation'}

    def _sku_price(mapping, sid):
        if not sid: return 0
        try:    s = mapping.get(int(sid))
        except: return 0
        return float(getattr(s, phase, 0) or 0) if s else 0

    al_total  = 0
    bom_lines = []
    for row in pricing_rows:
        qty    = int(row.get('qty', 0) or 0)
        app_p  = _sku_price(app_map,   row.get('appliance_id'))
        sdw_p  = _sku_price(sdwan_map, row.get('sdwan_id'))
        seg_p  = _sku_price(seg_map,   row.get('seg_id'))
        site_t = app_p + sdw_p + seg_p
        line_t = qty * site_t
        al_total += line_t
        if qty > 0:
            bom_lines.append((row.get('label', '—'), qty,
                              app_p, sdw_p, seg_p, site_t, line_t))

    support_amt = round(al_total * support_pct / 100)
    margin_amt  = round(al_total * margin_pct  / 100)
    annual_rec  = al_total + support_amt + margin_amt
    yr1_total   = annual_rec + services_amt
    _sm_mult    = 1 + (support_pct / 100) + (margin_pct / 100)

    tco_catalog = {e.id: float(e.annual_cost or 0)
                   for e in TCOEntry.query.filter_by(active=True).all()}
    TCO_CATS = ['fw_ns','sdwan','mpls','fw_ew','iot_ot','nac','l3sw','pam']
    legacy_hw = 0.0
    for row in tco_rows:
        qty = int(row.get('qty', 0) or 0)
        for cat in TCO_CATS:
            eid = row.get(f'{cat}_id')
            if eid:
                try:    legacy_hw += tco_catalog.get(int(eid), 0) * qty
                except: pass

    legacy_total = legacy_hw + fte_count * fte_cost + breach_cost
    annual_sav   = legacy_total - acv
    yr3_sav      = annual_sav * 3
    roi_pct      = round(annual_sav / legacy_total * 100) if legacy_total > 0 else 0
    payback_mo   = round((acv / annual_sav) * 12) if annual_sav > 0 else None

    prs, LY = _open_template()
    W = prs.slide_width
    H = prs.slide_height

    # ══════════════════════════════════════════════════════════════════════
    # SLIDE 1 — Cover
    # ══════════════════════════════════════════════════════════════════════
    s1 = prs.slides.add_slide(LY['cover'])
    _box(s1, 'Zero Trust Branch',
         Emu(388_620), Emu(2_000_000), Emu(8_229_600), Emu(457_200),
         sz=12, bold=True, color=CYAN)
    _box(s1, user_resp.customer_name or 'Customer',
         Emu(388_620), Emu(2_500_000), Emu(9_144_000), Emu(1_000_000),
         sz=36, bold=True, color=WHITE)
    meta = []
    if user_resp.se_name:      meta.append(user_resp.se_name)
    if user_resp.completed_at: meta.append(user_resp.completed_at.strftime('%B %d, %Y'))
    if meta:
        _box(s1, '     '.join(meta),
             Emu(388_620), Emu(3_600_000), Emu(9_144_000), Emu(457_200),
             sz=14, color=WHITE)

    # ══════════════════════════════════════════════════════════════════════
    # SLIDE 2 — Value Drivers
    # ══════════════════════════════════════════════════════════════════════
    s2 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s2, 'Value Drivers', 'Current State → Future State with Zscaler ZTB')

    drivers = [s for s in section_outcomes if s.get('format') == 'value_driver']
    if not drivers:
        _box(s2, 'No value drivers configured for this assessment.',
             Inches(0.5), Inches(2.0), Inches(11), Inches(0.5),
             sz=14, color=MUTED, italic=True)
    else:
        col_x = [Emu(370_819),  Emu(2_514_600), Emu(7_131_600)]
        col_w = [Emu(2_000_000), Emu(4_500_000), Emu(4_500_000)]
        HDR_Y = Emu(1_158_750)
        HDR_H = Emu(347_472)
        hdrs  = ['DRIVER', 'CURRENT STATE', 'FUTURE STATE (ZTB)']
        for i, h in enumerate(hdrs):
            _rect(s2, col_x[i], HDR_Y, col_w[i] - Emu(45_720), HDR_H, CARD)
            _box(s2, h, col_x[i] + Emu(63_500), HDR_Y,
                 col_w[i], HDR_H, sz=9, bold=True, color=CYAN)
        TABLE_TOP    = int(HDR_Y) + int(HDR_H)
        TABLE_BOTTOM = int(Inches(7.1))
        AVAILABLE    = TABLE_BOTTOM - TABLE_TOP
        n            = min(len(drivers), 7)
        def _est_lines(drv):
            def count(lines):
                total = 0
                for ln in lines:
                    words = len((ln or '').split())
                    total += max(1, -(-words // 10))
                return total
            return max(count(drv.get('current_state_lines', [])),
                       count(drv.get('future_state_lines',  [])), 1)
        weights     = [_est_lines(d) for d in drivers[:n]]
        total_w     = sum(weights)
        MIN_H       = int(Inches(0.7))
        raw_heights = [max(MIN_H, (w / total_w) * AVAILABLE) for w in weights]
        leftover    = AVAILABLE - sum(raw_heights)
        if leftover > 0:
            bonus = leftover / n
            raw_heights = [rh + bonus for rh in raw_heights]
        y_cursor = TABLE_TOP
        for ri, drv in enumerate(drivers[:n]):
            rh  = int(raw_heights[ri])
            bg_ = CARD if ri % 2 == 0 else CARD2
            PAD = int(Emu(54_864))
            for i in range(3):
                _rect(s2, col_x[i], y_cursor, col_w[i] - Emu(45_720),
                      rh - int(Emu(27_432)), bg_)
            _box(s2, drv.get('label', ''),
                 col_x[0] + Emu(91_440), y_cursor + PAD, col_w[0] - Emu(137_160), rh,
                 sz=11, bold=True, color=CYAN)
            _box(s2, '\n'.join(drv.get('current_state_lines', [])),
                 col_x[1] + Emu(73_152), y_cursor + PAD, col_w[1] - Emu(137_160), rh,
                 sz=11, color=WHITE)
            _box(s2, '\n'.join(drv.get('future_state_lines', [])),
                 col_x[2] + Emu(73_152), y_cursor + PAD, col_w[2] - Emu(137_160), rh,
                 sz=11, color=WHITE)
            y_cursor += rh

    # ══════════════════════════════════════════════════════════════════════
    # SLIDE 3 — Pricing Executive Summary
    # ══════════════════════════════════════════════════════════════════════
    s3 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s3, 'Pricing Summary', 'Executive Overview')

    COL_X = [Emu(370_819),  Emu(5_303_520), Emu(8_601_120)]
    COL_W = [Emu(4_750_000), Emu(3_100_000), Emu(3_200_000)]

    HDR_Y = Emu(1_158_750); HDR_H = Emu(383_731)
    for i, (h, al) in enumerate([('SITE LABEL', PP_ALIGN.LEFT),
                                  ('COST PER SITE', PP_ALIGN.RIGHT),
                                  ('LINE TOTAL', PP_ALIGN.RIGHT)]):
        _rect(s3, COL_X[i], HDR_Y, COL_W[i] - Emu(45_720), HDR_H, CARD)
        _box(s3, h, COL_X[i] + Emu(127_000), HDR_Y + Emu(45_720),
             COL_W[i] - Emu(165_100), HDR_H, sz=10, bold=True, color=CYAN, align=al)

    active = [
        (lbl, round(site_t * _sm_mult), round(qty * site_t * _sm_mult))
        for lbl, qty, app_p, sdw_p, seg_p, site_t, line_t in bom_lines
    ]

    TABLE_TOP    = int(HDR_Y) + int(HDR_H)
    TABLE_BOTTOM = int(Inches(3.8))
    n_rows       = max(len(active), 1)
    ROW_H        = max(int(Inches(0.45)), int((TABLE_BOTTOM - TABLE_TOP) / n_rows))

    y_cur = TABLE_TOP
    for ri, (lbl, cost_per_site, line_total) in enumerate(active):
        bg_  = CARD if ri % 2 == 0 else CARD2
        PAD  = int(Emu(109_728))
        for i in range(3):
            _rect(s3, COL_X[i], y_cur, COL_W[i] - Emu(45_720),
                  ROW_H - int(Emu(27_432)), bg_)
        _box(s3, lbl,
             COL_X[0] + Emu(127_000), y_cur + PAD,
             COL_W[0] - Emu(182_880), ROW_H, sz=14, bold=True, color=CYAN)
        _box(s3, _fmt_exact(cost_per_site),
             COL_X[1] + Emu(45_720), y_cur + PAD,
             COL_W[1] - Emu(91_440), ROW_H, sz=14, color=WHITE, align=PP_ALIGN.RIGHT)
        _box(s3, _fmt_exact(line_total),
             COL_X[2] + Emu(45_720), y_cur + PAD,
             COL_W[2] - Emu(91_440), ROW_H, sz=14, bold=True, color=WHITE, align=PP_ALIGN.RIGHT)
        y_cur += ROW_H

    GT_Y = y_cur + int(Emu(109_728))
    _rect(s3, Emu(370_819), GT_Y, Emu(11_430_381), Emu(640_080), CARD)
    _rect(s3, Emu(370_819), GT_Y, Emu(73_152),     Emu(640_080), BLUE)
    _box(s3, 'GRAND TOTAL',
         Emu(548_640), GT_Y + Emu(109_728), Emu(4_572_000), Emu(457_200),
         sz=16, bold=True, color=WHITE)
    _box(s3, _fmt_exact(annual_rec),
         Emu(5_303_520), GT_Y + Emu(73_152), Emu(6_400_000), Emu(548_640),
         sz=22, bold=True, color=CYAN, align=PP_ALIGN.RIGHT)

    LIC_Y = GT_Y + int(Emu(1_188_720))
    _box(s3, 'LICENSES INCLUDED',
         Emu(370_819), LIC_Y - Emu(256_032), Emu(10_972_800), Emu(237_744),
         sz=9, bold=True, color=MUTED)

    has_sdwan  = any(row.get('sdwan_id') for row in pricing_rows)
    has_seg    = any(row.get('seg_id')   for row in pricing_rows)
    lic_labels = []
    if has_sdwan: lic_labels.append('SD-WAN Licenses')
    if has_seg:   lic_labels.append('Segmentation Licenses')
    if not lic_labels: lic_labels = ['Appliance only — no software licenses']

    n_pills = len(lic_labels)
    pill_w  = int((int(Emu(11_430_381)) - int(Emu(137_160)) * (n_pills - 1)) / n_pills)
    pill_x  = int(Emu(370_819))
    for lic in lic_labels:
        _rect(s3, pill_x, LIC_Y, pill_w, int(Emu(502_920)), CARD)
        _rect(s3, pill_x, LIC_Y, int(Emu(64_008)), int(Emu(502_920)), BLUE)
        _box(s3, lic, pill_x + int(Emu(146_304)), LIC_Y + int(Emu(91_440)),
             pill_w - int(Emu(219_456)), int(Emu(383_731)), sz=13, bold=True, color=WHITE)
        pill_x += pill_w + int(Emu(137_160))

    _box(s3, f'Phase: {phase.upper()}  ·  Annual Recurring: {_fmt(annual_rec)}',
         Emu(370_819), LIC_Y + Emu(594_360), Emu(11_430_381), Emu(274_320),
         sz=9, color=MUTED, italic=True, align=PP_ALIGN.RIGHT)

    # ══════════════════════════════════════════════════════════════════════
    # SLIDE 4 — TCO KPIs + Bar Chart
    # ══════════════════════════════════════════════════════════════════════
    s4 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s4, 'TCO Analysis', 'Legacy Infrastructure vs. Zscaler ZTB')

    kpis = [
        ('Legacy Annual Spend', _fmt(legacy_total), RED,   'HW + FTE + Breach Risk'),
        ('Annual Savings',      _fmt(annual_sav),   GREEN, 'Legacy − Zscaler ACV'),
        ('3-Year Savings',      _fmt(yr3_sav),      GREEN, '3× Annual Savings'),
        ('ROI',                 f'{roi_pct}%',       AMBER,
         f'Payback: ~{payback_mo} months' if payback_mo else 'Set ACV above'),
    ]
    kw = Emu(2_743_200); kh = Emu(1_417_320); kg = Emu(164_592)
    for i, (title, val, col, sub) in enumerate(kpis):
        kx = int(Emu(370_819)) + i * (int(kw) + int(kg))
        _rect(s4, kx, Emu(868_680), kw, kh, CARD)
        _rect(s4, kx, Emu(868_680), Emu(64_008), kh, col)
        _box(s4, title, kx + Emu(137_160), Emu(978_408), kw, Emu(320_040),
             sz=8, bold=True, color=MUTED)
        _box(s4, val,   kx + Emu(91_440),  Emu(1_280_160), kw - Emu(91_440), Emu(548_640),
             sz=26, bold=True, color=col)
        if sub:
            _box(s4, sub, kx + Emu(137_160), Emu(1_826_640), kw, Emu(274_320),
                 sz=8, color=MUTED, italic=True)

    fig, ax = plt.subplots(figsize=(12.5, 3.5), facecolor='#001744')
    ax.set_facecolor('#001744')
    x_pos = [0, 1, 2]
    ax.bar([x - 0.22 for x in x_pos], [legacy_total] * 3, 0.38,
           color='#ED1951', label='Legacy Annual Cost', zorder=3)
    ax.bar([x + 0.22 for x in x_pos], [acv] * 3, 0.38,
           color='#12D3FF', label='Zscaler ACV', zorder=3)
    ax.set_xticks(x_pos)
    ax.set_xticklabels(['Year 1', 'Year 2', 'Year 3'],
                       color='#FFFFFF', fontsize=11, fontfamily='DejaVu Sans')
    ax.yaxis.set_major_formatter(mticker.FuncFormatter(
        lambda v, _: f'${v/1e6:.1f}M' if v >= 1e6 else f'${v/1e3:.0f}K'))
    ax.tick_params(axis='y', colors='#FFFFFF', labelsize=10)
    for sp in ax.spines.values():
        sp.set_visible(False)
    ax.yaxis.grid(True, color='#002466', linewidth=0.7, zorder=0)
    ax.set_axisbelow(True)
    ax.legend(facecolor='#001744', edgecolor='#246CF7',
              labelcolor='#FFFFFF', fontsize=10, loc='upper right')
    plt.tight_layout(pad=0.3)
    chart_buf = io.BytesIO()
    fig.savefig(chart_buf, format='png', dpi=130,
                facecolor='#001744', bbox_inches='tight')
    plt.close(fig)
    chart_buf.seek(0)
    s4.shapes.add_picture(chart_buf, Emu(320_040), Emu(2_400_000),
                          Emu(11_521_440), Emu(3_700_000))

    # ══════════════════════════════════════════════════════════════════════
    # SLIDE 5 — Thank You
    # ══════════════════════════════════════════════════════════════════════
    s5 = prs.slides.add_slide(LY['cover'])
    _box(s5, 'Thank You',
         Emu(388_620), Emu(2_500_000), Emu(9_144_000), Emu(914_400),
         sz=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    line = user_resp.customer_name or ''
    if user_resp.se_name: line += f'  ·  {user_resp.se_name}'
    _box(s5, line,
         Emu(388_620), Emu(3_500_000), Emu(9_144_000), Emu(457_200),
         sz=18, color=CYAN, align=PP_ALIGN.CENTER)

    out   = io.BytesIO()
    prs.save(out)
    out.seek(0)
    safe  = (user_resp.customer_name or 'Customer').replace(' ', '_').replace('/', '_')
    fname = f"ZTB_{safe}_{datetime.utcnow().strftime('%Y%m%d')}.pptx"
    return send_file(
        out,
        mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        as_attachment=True,
        download_name=fname
    )


# ══════════════════════════════════════════════════════════════════════════════
# POV DECK EXPORT
# ══════════════════════════════════════════════════════════════════════════════

@export_ppt_bp.route('/user/<customer_slug>/export-pov-deck', methods=['POST'])
def export_pov_deck(customer_slug):
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    user_resp = _find_by_slug(customer_slug)
    if not user_resp:
        return jsonify({'error': 'Not found'}), 404

    NAVY  = RGBColor(0x00, 0x17, 0x44)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    BLUE  = RGBColor(0x24, 0x6C, 0xF7)
    CYAN  = RGBColor(0x12, 0xD3, 0xFF)
    GREEN = RGBColor(0x6B, 0xFF, 0xB3)
    AMBER = RGBColor(0xFF, 0x93, 0x00)
    MUTED = RGBColor(0xA0, 0xBC, 0xD8)
    CARD  = RGBColor(0x00, 0x23, 0x6B)
    CARD2 = RGBColor(0x00, 0x1F, 0x5E)

    section_outcomes = evaluate_sections(user_resp.answers or {})

    prs, LY = _open_template()
    W = prs.slide_width
    H = prs.slide_height

    pov_raw   = user_resp.raw_responses or {}
    prepov    = pov_raw.get('prepov_data', {})
    checks    = prepov.get('checks', {})
    notes_map = prepov.get('notes', {})
    tl_rows   = prepov.get('pov_timeline_v2', [])

    from routes.user_routes import _collect
    results   = user_resp.results or {}
    asset_ids = results.get('asset_ids', [])
    all_assets = _collect('assets', asset_ids)
    success_criteria = [
        a for a in all_assets
        if (a.get('asset_type') or '').lower() in ('success criteria', 'success_criteria')
    ]

    CHECKLIST_ITEMS = [
        ('arch_workshops',   'Architecture & Design Workshops'),
        ('stakeholders',     'Stakeholders Defined'),
        ('budget_owners',    'Budget Owners Defined'),
        ('costs_defined',    'Estimated Costs & Commercials Defined'),
        ('test_cases',       'POV Test Cases, Requirements & Scope Defined'),
        ('tech_feasibility', 'Technical Feasibility Review'),
        ('logistics',        'Logistics, Timeline & Roles Defined'),
        ('signoffs',         'POV Sign-Offs Completed by all Parties'),
    ]

    STAKEHOLDER_DEFS = [
        ('zs_champion',  'Zscaler Champion',         False),
        ('exec_sponsor', 'Executive Sponsor',         False),
        ('net_lead',     'Networking Lead',           False),
        ('sec_lead',     'Security Lead',             False),
        ('zs_ae',        'Zscaler Account Executive', True),
        ('zs_se',        'Zscaler Sales Engineer',    True),
    ]

    # ── SLIDE 1 — Cover ───────────────────────────────────────────────────
    s1 = prs.slides.add_slide(LY['cover'])
    _box(s1, 'PROOF OF VALUE — POV DECK',
         Emu(388_620), Emu(1_800_000), Emu(9_144_000), Emu(457_200),
         sz=12, bold=True, color=CYAN)
    _box(s1, user_resp.customer_name or 'Customer',
         Emu(388_620), Emu(2_300_000), Emu(9_144_000), Emu(1_000_000),
         sz=36, bold=True, color=WHITE)
    meta = []
    if user_resp.se_name:      meta.append(user_resp.se_name)
    if user_resp.completed_at: meta.append(user_resp.completed_at.strftime('%B %d, %Y'))
    if meta:
        _box(s1, '     '.join(meta),
             Emu(388_620), Emu(3_400_000), Emu(9_144_000), Emu(457_200),
             sz=14, color=WHITE)

    # ── SLIDE 2 — Pre-POV Checklist ───────────────────────────────────────
    s2 = prs.slides.add_slide(LY['title_sub'])
    done_count = sum(1 for iid, _ in CHECKLIST_ITEMS if checks.get(iid))
    _slide_header(s2, 'Pre-POV Checklist',
                  f'Completion: {done_count} of {len(CHECKLIST_ITEMS)} items')

    TABLE_TOP    = int(Emu(1_158_750)) + int(Emu(347_472))
    TABLE_BOTTOM = int(Inches(6.7))
    ROW_H = max(int((TABLE_BOTTOM - TABLE_TOP) / len(CHECKLIST_ITEMS)), int(Inches(0.62)))
    COL_ST_X = Emu(370_819);   COL_ST_W = Emu(502_920)
    COL_TT_X = Emu(914_399);   COL_TT_W = Emu(5_121_960)
    COL_NT_X = Emu(6_127_679); COL_NT_W = Emu(5_850_721)
    HDR_Y = TABLE_TOP; HDR_H = int(Emu(347_472))
    for cx, cw, lbl in [(COL_ST_X, COL_ST_W, ''),
                         (COL_TT_X, COL_TT_W, 'CHECKLIST ITEM'),
                         (COL_NT_X, COL_NT_W, 'NOTES')]:
        _rect(s2, cx, HDR_Y, cw, HDR_H, CARD)
        if lbl:
            _box(s2, lbl, cx + Emu(73_152), HDR_Y, cw, HDR_H,
                 sz=9, bold=True, color=GREEN)
    y = HDR_Y + HDR_H
    for i, (item_id, item_title) in enumerate(CHECKLIST_ITEMS):
        checked = bool(checks.get(item_id))
        bg_ = CARD if i % 2 == 0 else CARD2
        PAD = int(Emu(91_440))
        for cx, cw in [(COL_ST_X, COL_ST_W), (COL_TT_X, COL_TT_W), (COL_NT_X, COL_NT_W)]:
            _rect(s2, cx, y, cw, ROW_H - int(Emu(27_432)), bg_)
        _box(s2, '✅' if checked else '⬜',
             COL_ST_X + Emu(36_576), y + PAD, COL_ST_W, ROW_H,
             sz=16, color=GREEN if checked else MUTED, align=PP_ALIGN.CENTER)
        _box(s2, item_title,
             COL_TT_X + Emu(73_152), y + PAD, COL_TT_W - Emu(109_728), ROW_H,
             sz=11, bold=(not checked), color=MUTED if checked else WHITE)
        note_text = notes_map.get(item_id, '')
        if note_text:
            _box(s2, note_text,
                 COL_NT_X + Emu(73_152), y + PAD, COL_NT_W - Emu(109_728), ROW_H,
                 sz=10, color=MUTED, italic=True)
        y += ROW_H

    # ── SLIDE 3 — Key Stakeholders ────────────────────────────────────────
    s3 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s3, 'Key Stakeholders', 'POV Team & Sign-Off Contacts')

    stk_data = prepov.get('stakeholders', {})

    COL_ROLE_X = Emu(370_819);   COL_ROLE_W = Emu(3_474_720)
    COL_NAME_X = Emu(3_936_339); COL_NAME_W = Emu(4_846_320)
    COL_DATE_X = Emu(8_874_099); COL_DATE_W = Emu(2_970_381)

    HDR_Y = int(Emu(1_158_750)) + int(Emu(347_472)); HDR_H = int(Emu(347_472))
    for cx, cw, lbl, al in [
        (COL_ROLE_X, COL_ROLE_W, 'ROLE',          PP_ALIGN.LEFT),
        (COL_NAME_X, COL_NAME_W, 'NAME & TITLE',   PP_ALIGN.LEFT),
        (COL_DATE_X, COL_DATE_W, 'SIGN-OFF DATE',  PP_ALIGN.RIGHT),
    ]:
        _rect(s3, cx, HDR_Y, cw, HDR_H, CARD)
        _box(s3, lbl, cx + Emu(73_152), HDR_Y, cw, HDR_H,
             sz=9, bold=True, color=CYAN, align=al)

    TABLE_TOP    = HDR_Y + HDR_H
    TABLE_BOTTOM = int(Inches(6.7))
    n_stk = len(STAKEHOLDER_DEFS)
    ROW_H = max(int((TABLE_BOTTOM - TABLE_TOP) / n_stk), int(Inches(0.62)))
    zscaler_div_done = False
    y = TABLE_TOP

    for i, (stk_id, stk_role, is_zscaler) in enumerate(STAKEHOLDER_DEFS):
        if is_zscaler and not zscaler_div_done:
            total_w = int(COL_ROLE_W) + int(COL_NAME_W) + int(Emu(91_440)) + int(COL_DATE_W)
            _rect(s3, COL_ROLE_X, y, total_w, int(Emu(256_032)), CARD)
            _box(s3, 'ZSCALER TEAM', COL_ROLE_X + Emu(109_728), y + Emu(36_576),
                 Emu(3_657_600), int(Emu(256_032)), sz=7, bold=True, color=CYAN)
            y += int(Emu(256_032))
            zscaler_div_done = True

        bg_ = CARD if i % 2 == 0 else CARD2
        PAD = int(Emu(109_728))
        for cx, cw in [(COL_ROLE_X, COL_ROLE_W), (COL_NAME_X, COL_NAME_W),
                       (COL_DATE_X, COL_DATE_W)]:
            _rect(s3, cx, y, cw, ROW_H - int(Emu(27_432)), bg_)
        _box(s3, stk_role,
             COL_ROLE_X + Emu(91_440), y + PAD, COL_ROLE_W - Emu(137_160), ROW_H,
             sz=11, bold=True, color=WHITE)
        person   = stk_data.get(stk_id, {})
        name_val = person.get('name', '') if isinstance(person, dict) else ''
        date_val = person.get('date', '') if isinstance(person, dict) else ''
        if name_val:
            _box(s3, name_val,
                 COL_NAME_X + Emu(91_440), y + PAD, COL_NAME_W - Emu(137_160), ROW_H,
                 sz=11, color=WHITE)
        else:
            _box(s3, '—',
                 COL_NAME_X + Emu(91_440), y + PAD, COL_NAME_W - Emu(137_160), ROW_H,
                 sz=11, color=MUTED, italic=True)
        if not is_zscaler:
            if date_val:
                _box(s3, date_val,
                     COL_DATE_X + Emu(54_864), y + PAD, COL_DATE_W - Emu(91_440), ROW_H,
                     sz=11, bold=True, color=GREEN, align=PP_ALIGN.RIGHT)
            else:
                _box(s3, 'Pending',
                     COL_DATE_X + Emu(54_864), y + PAD, COL_DATE_W - Emu(91_440), ROW_H,
                     sz=10, color=AMBER, italic=True, align=PP_ALIGN.RIGHT)
        y += ROW_H

    # ── SLIDE 4 — POV Timeline ────────────────────────────────────────────
    s4 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s4, 'POV Timeline', 'Milestone schedule for the Proof of Value')

    if not tl_rows:
        _box(s4, 'No timeline milestones have been defined yet.',
             Inches(0.5), Inches(2.0), Inches(12), Inches(0.5),
             sz=14, color=MUTED, italic=True)
    else:
        TABLE_TOP    = int(Emu(1_158_750)) + int(Emu(347_472))
        TABLE_BOTTOM = int(Inches(6.7))
        n = min(len(tl_rows), 14)
        ROW_H = max(int((TABLE_BOTTOM - TABLE_TOP) / n), int(Inches(0.38)))
        COL_NUM_X = Emu(370_819);   COL_NUM_W = Emu(457_200)
        COL_MIL_X = Emu(868_019);   COL_MIL_W = Emu(7_940_160)
        COL_DAT_X = Emu(8_899_419); COL_DAT_W = Emu(2_945_061)
        HDR_Y = TABLE_TOP; HDR_H = int(Emu(347_472))
        for cx, cw, lbl, al in [
            (COL_NUM_X, COL_NUM_W, '#',            PP_ALIGN.CENTER),
            (COL_MIL_X, COL_MIL_W, 'MILESTONE',   PP_ALIGN.LEFT),
            (COL_DAT_X, COL_DAT_W, 'TARGET DATE',  PP_ALIGN.RIGHT),
        ]:
            _rect(s4, cx, HDR_Y, cw, HDR_H, CARD)
            _box(s4, lbl, cx + Emu(54_864), HDR_Y, cw, HDR_H,
                 sz=9, bold=True, color=AMBER, align=al)
        y = HDR_Y + HDR_H
        for i, row in enumerate(tl_rows[:n]):
            bg_ = CARD if i % 2 == 0 else CARD2
            PAD = int(Emu(73_152))
            for cx, cw in [(COL_NUM_X, COL_NUM_W), (COL_MIL_X, COL_MIL_W),
                           (COL_DAT_X, COL_DAT_W)]:
                _rect(s4, cx, y, cw, ROW_H - int(Emu(18_288)), bg_)
            _box(s4, str(i + 1), COL_NUM_X + Emu(54_864), y + PAD, COL_NUM_W, ROW_H,
                 sz=10, bold=True, color=MUTED, align=PP_ALIGN.CENTER)
            _box(s4, row.get('label', ''), COL_MIL_X + Emu(73_152), y + PAD,
                 COL_MIL_W - Emu(109_728), ROW_H, sz=11, color=WHITE)
            if row.get('date', ''):
                _box(s4, row['date'], COL_DAT_X + Emu(54_864), y + PAD,
                     COL_DAT_W - Emu(91_440), ROW_H, sz=11, bold=True,
                     color=CYAN, align=PP_ALIGN.RIGHT)
            y += ROW_H

    # ── SLIDE 5 — Value Drivers ───────────────────────────────────────────
    s5 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s5, 'Value Drivers', 'Current State → Future State with Zscaler ZTB')

    drivers = [s for s in section_outcomes if s.get('format') == 'value_driver']
    if not drivers:
        _box(s5, 'No value drivers configured for this assessment.',
             Inches(0.5), Inches(2.0), Inches(11), Inches(0.5),
             sz=14, color=MUTED, italic=True)
    else:
        col_x = [Emu(370_819),  Emu(2_514_600), Emu(7_131_600)]
        col_w = [Emu(2_000_000), Emu(4_500_000), Emu(4_500_000)]
        HDR_Y = Emu(1_158_750)
        HDR_H = Emu(347_472)
        hdrs  = ['DRIVER', 'CURRENT STATE', 'FUTURE STATE (ZTB)']
        for i, h in enumerate(hdrs):
            _rect(s5, col_x[i], HDR_Y, col_w[i] - Emu(45_720), HDR_H, CARD)
            _box(s5, h, col_x[i] + Emu(63_500), HDR_Y,
                 col_w[i], HDR_H, sz=9, bold=True, color=CYAN)
        TABLE_TOP    = int(HDR_Y) + int(HDR_H)
        TABLE_BOTTOM = int(Inches(7.1))
        AVAILABLE    = TABLE_BOTTOM - TABLE_TOP
        n            = min(len(drivers), 7)
        def _est_lines(drv):
            def count(lines):
                total = 0
                for ln in lines:
                    words = len((ln or '').split())
                    total += max(1, -(-words // 10))
                return total
            return max(count(drv.get('current_state_lines', [])),
                       count(drv.get('future_state_lines',  [])), 1)
        weights     = [_est_lines(d) for d in drivers[:n]]
        total_w     = sum(weights)
        MIN_H       = int(Inches(0.7))
        raw_heights = [max(MIN_H, (w / total_w) * AVAILABLE) for w in weights]
        leftover    = AVAILABLE - sum(raw_heights)
        if leftover > 0:
            bonus = leftover / n
            raw_heights = [rh + bonus for rh in raw_heights]
        y_cursor = TABLE_TOP
        for ri, drv in enumerate(drivers[:n]):
            rh  = int(raw_heights[ri])
            bg_ = CARD if ri % 2 == 0 else CARD2
            PAD = int(Emu(54_864))
            for i in range(3):
                _rect(s5, col_x[i], y_cursor, col_w[i] - Emu(45_720),
                      rh - int(Emu(27_432)), bg_)
            _box(s5, drv.get('label', ''),
                 col_x[0] + Emu(91_440), y_cursor + PAD, col_w[0] - Emu(137_160), rh,
                 sz=11, bold=True, color=CYAN)
            _box(s5, '\n'.join(drv.get('current_state_lines', [])),
                 col_x[1] + Emu(73_152), y_cursor + PAD, col_w[1] - Emu(137_160), rh,
                 sz=11, color=WHITE)
            _box(s5, '\n'.join(drv.get('future_state_lines', [])),
                 col_x[2] + Emu(73_152), y_cursor + PAD, col_w[2] - Emu(137_160), rh,
                 sz=11, color=WHITE)
            y_cursor += rh

    # ── SLIDE 6 — Success Criteria ────────────────────────────────────────
    s6 = prs.slides.add_slide(LY['title_sub'])
    _slide_header(s6, 'POV Success Criteria', 'Defined test cases and acceptance criteria')

    if not success_criteria:
        _box(s6, 'No success criteria have been tagged for this assessment.',
             Inches(0.5), Inches(2.0), Inches(12), Inches(0.5),
             sz=14, color=MUTED, italic=True)
    else:
        TABLE_TOP    = int(Emu(1_158_750)) + int(Emu(347_472))
        TABLE_BOTTOM = int(Inches(6.7))
        n         = min(len(success_criteria), 8)
        ROW_H     = max(int(Inches(0.55)), int((TABLE_BOTTOM - TABLE_TOP) / n))
        COL_NUM_X = Emu(370_819);   COL_NUM_W = Emu(457_200)
        COL_TIT_X = Emu(868_019);   COL_TIT_W = Emu(5_303_520)
        COL_DSC_X = Emu(6_217_399); COL_DSC_W = Emu(5_713_981)
        HDR_Y = TABLE_TOP; HDR_H = int(Emu(347_472))
        for cx, cw, lbl in [(COL_NUM_X, COL_NUM_W, '#'),
                             (COL_TIT_X, COL_TIT_W, 'SUCCESS CRITERIA'),
                             (COL_DSC_X, COL_DSC_W, 'DESCRIPTION')]:
            _rect(s6, cx, HDR_Y, cw, HDR_H, CARD)
            _box(s6, lbl, cx + Emu(73_152), HDR_Y, cw, HDR_H,
                 sz=9, bold=True, color=BLUE)
        y = HDR_Y + HDR_H
        for i, sc in enumerate(success_criteria[:n]):
            bg_ = CARD if i % 2 == 0 else CARD2
            PAD = int(Emu(91_440))
            for cx, cw in [(COL_NUM_X, COL_NUM_W), (COL_TIT_X, COL_TIT_W),
                           (COL_DSC_X, COL_DSC_W)]:
                _rect(s6, cx, y, cw, ROW_H - int(Emu(27_432)), bg_)
            _box(s6, str(i + 1), COL_NUM_X + Emu(73_152), y + PAD, COL_NUM_W, ROW_H,
                 sz=11, bold=True, color=CYAN, align=PP_ALIGN.CENTER)
            title = sc.get('asset_name') or sc.get('title') or f'Criteria {i+1}'
            _box(s6, title, COL_TIT_X + Emu(73_152), y + PAD,
                 COL_TIT_W - Emu(109_728), ROW_H, sz=11, bold=True, color=WHITE)
            desc = sc.get('description') or ''
            if desc:
                _box(s6, desc, COL_DSC_X + Emu(73_152), y + PAD,
                     COL_DSC_W - Emu(109_728), ROW_H, sz=10, color=MUTED)
            y += ROW_H

    # ── SLIDE 7 — Deployment Pattern ─────────────────────────────────────
    DEPLOY_PATH   = '/home/ubuntu/ztb-super-app/static/assets/pov_deployment_patterns.pptx'
    DIAGRAM_PATHS = [
        '/home/ubuntu/ztb-super-app/static/assets/diagram_type1.png',
        '/home/ubuntu/ztb-super-app/static/assets/diagram_type2.png',
        '/home/ubuntu/ztb-super-app/static/assets/diagram_type3.png',
    ]
    try:
        deploy_outcome = next(
            (s for s in section_outcomes if s.get('name', '').strip().lower() == 'deployment type'),
            None
        )
        outcome_str = (deploy_outcome.get('outcome', '') or '').strip().upper() if deploy_outcome else ''
        if 'TYPE 3' in outcome_str:
            slide_idx = 2
        elif 'TYPE 2' in outcome_str:
            slide_idx = 1
        else:
            slide_idx = 0
        _append_slides_from_file(prs, DEPLOY_PATH, slide_index=slide_idx)
        _inject_diagram_png(prs.slides[-1], DIAGRAM_PATHS[slide_idx])
    except Exception:
        import traceback; traceback.print_exc()

    # ── SLIDES 8, 9, 10 — Universal POV slides ───────────────────────────
    UNIVERSAL_PATH = '/home/ubuntu/ztb-super-app/static/assets/pov_deck_universal_slides.pptx'
    try:
        _append_slides_from_file(prs, UNIVERSAL_PATH)
        slide10 = prs.slides[-1]
        for shape in slide10.shapes:
            if shape.top is not None and shape.top < int(Emu(300_000)):
                shape.top = shape.top + int(Emu(400_000))
    except Exception:
        pass

    # ── SLIDE 11 — Thank You ──────────────────────────────────────────────
    s_ty = prs.slides.add_slide(LY['cover'])
    _box(s_ty, 'Thank You',
         Emu(388_620), Emu(2_500_000), Emu(9_144_000), Emu(914_400),
         sz=40, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    line = user_resp.customer_name or ''
    if user_resp.se_name: line += f'  ·  {user_resp.se_name}'
    _box(s_ty, line,
         Emu(388_620), Emu(3_500_000), Emu(9_144_000), Emu(457_200),
         sz=18, color=CYAN, align=PP_ALIGN.CENTER)

    out   = io.BytesIO()
    prs.save(out)
    out.seek(0)
    safe  = (user_resp.customer_name or 'Customer').replace(' ', '_').replace('/', '_')
    fname = f"ZTB_POV_Deck_{safe}_{datetime.utcnow().strftime('%Y%m%d')}.pptx"
    return send_file(
        out,
        mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        as_attachment=True,
        download_name=fname
    )
