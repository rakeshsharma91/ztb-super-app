from flask import Blueprint, render_template, request, jsonify, session, redirect, url_for, send_file
from models import db, PricingSKU, UserResponse
import re, os, json
from datetime import datetime, timezone
from werkzeug.utils import secure_filename

pricing_bp      = Blueprint('pricing', __name__)
pricing_user_bp = Blueprint('pricing_user', __name__)

TEMPLATE_DIR  = os.path.join(os.path.dirname(__file__), '..', 'static', 'pptx_templates')
TEMPLATE_PATH = os.path.join(TEMPLATE_DIR, 'executive_template.pptx')
TEMPLATE_META = os.path.join(TEMPLATE_DIR, 'template_meta.json')

os.makedirs(TEMPLATE_DIR, exist_ok=True)

VALID_SIZES = ('Small', 'Medium', 'Large', 'XL')

# ── Auth helper ───────────────────────────────────────────────────────────────
def admin_required():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin.login'))
    return None


# ══════════════════════════════════════════════════════════════════════════════
# ADMIN — Pricing Grid page
# ══════════════════════════════════════════════════════════════════════════════

@pricing_bp.route('/admin/pricing/')
def pricing_grid():
    redir = admin_required()
    if redir:
        return redir
    return render_template('pricing.html')


# ── GET all SKUs ──────────────────────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/skus', methods=['GET'])
def get_skus():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    skus = PricingSKU.query.order_by(PricingSKU.display_order, PricingSKU.id).all()
    return jsonify({'success': True, 'skus': [s.to_dict() for s in skus]})


# ── CREATE SKU ────────────────────────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/skus', methods=['POST'])
def create_sku():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    data = request.get_json() or {}
    raw_size = data.get('size') or None
    sku = PricingSKU(
        sku_code      = data.get('sku_code', '').strip(),
        sku_name      = data.get('sku_name', '').strip(),
        category      = data.get('category', 'appliance'),
        size          = raw_size if raw_size in VALID_SIZES else None,
        cogs          = float(data.get('cogs', 0) or 0),
        list_price    = float(data.get('list_price', 0) or 0),
        budgetary     = float(data.get('budgetary', 0) or 0),
        standard      = float(data.get('standard', 0) or 0),
        aggressive    = float(data.get('aggressive', 0) or 0),
        is_ha         = bool(data.get('is_ha', False)),
        ha_parent_id  = int(data['ha_parent_id']) if data.get('ha_parent_id') else None,
        display_order = int(data.get('display_order', 99) or 99),
        active        = bool(data.get('active', True)),
    )
    db.session.add(sku)
    db.session.commit()
    return jsonify({'success': True, 'sku': sku.to_dict()})


# ── UPDATE SKU ────────────────────────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/skus/<int:sku_id>', methods=['PUT'])
def update_sku(sku_id):
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    sku = PricingSKU.query.get_or_404(sku_id)
    data = request.get_json() or {}

    if 'sku_code'      in data: sku.sku_code      = data['sku_code'].strip()
    if 'sku_name'      in data: sku.sku_name      = data['sku_name'].strip()
    if 'category'      in data: sku.category      = data['category']
    if 'size'          in data:
        raw_size = data['size'] or None
        sku.size = raw_size if raw_size in VALID_SIZES else None
    if 'cogs'          in data: sku.cogs          = float(data['cogs'] or 0)
    if 'list_price'    in data: sku.list_price    = float(data['list_price'] or 0)
    if 'budgetary'     in data: sku.budgetary     = float(data['budgetary'] or 0)
    if 'standard'      in data: sku.standard      = float(data['standard'] or 0)
    if 'aggressive'    in data: sku.aggressive    = float(data['aggressive'] or 0)
    if 'is_ha'         in data: sku.is_ha         = bool(data['is_ha'])
    if 'ha_parent_id'  in data: sku.ha_parent_id  = int(data['ha_parent_id']) if data['ha_parent_id'] else None
    if 'display_order' in data: sku.display_order = int(data['display_order'] or 0)
    if 'active'        in data: sku.active        = bool(data['active'])

    if data.get('sync_ha') and not sku.is_ha:
        ha_child = PricingSKU.query.filter_by(ha_parent_id=sku.id, is_ha=True).first()
        if ha_child:
            ha_child.list_price = sku.list_price * 2
            ha_child.budgetary  = sku.budgetary  * 2
            ha_child.standard   = sku.standard   * 2
            ha_child.aggressive = sku.aggressive * 2
            ha_child.cogs       = sku.cogs       * 2

    db.session.commit()
    return jsonify({'success': True, 'sku': sku.to_dict()})


# ── DELETE (deactivate) SKU ───────────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/skus/<int:sku_id>', methods=['DELETE'])
def delete_sku(sku_id):
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    sku = PricingSKU.query.get_or_404(sku_id)
    sku.active = False
    db.session.commit()
    return jsonify({'success': True})


# ── HARD DELETE SKU ───────────────────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/skus/<int:sku_id>/hard', methods=['DELETE'])
def hard_delete_sku(sku_id):
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    sku = PricingSKU.query.get_or_404(sku_id)
    db.session.delete(sku)
    db.session.commit()
    return jsonify({'success': True})


# ══════════════════════════════════════════════════════════════════════════════
# USER — Pricing tab on customer results page
# ══════════════════════════════════════════════════════════════════════════════

def _slug_to_response(customer_slug):
    """Return the UserResponse for a slug (same logic as user_routes)."""
    name = customer_slug.replace('-', ' ')
    resp = UserResponse.query.filter(
        db.func.lower(UserResponse.customer_name) == name.lower(),
        UserResponse.status == 'completed'
    ).order_by(UserResponse.completed_at.desc()).first()
    return resp


# ── GET pricing config (SKUs + saved rows + site_profile for this customer) ───
@pricing_user_bp.route('/user/<customer_slug>/pricing-config', methods=['GET'])
def pricing_config(customer_slug):
    skus = PricingSKU.query.filter_by(active=True).order_by(
        PricingSKU.display_order, PricingSKU.id
    ).all()

    appliances   = [s.to_dict() for s in skus if s.category == 'appliance']
    sdwan        = [s.to_dict() for s in skus if s.category == 'sdwan']
    segmentation = [s.to_dict() for s in skus if s.category == 'segmentation']

    saved         = {}
    site_profile  = []
    resp = _slug_to_response(customer_slug)
    if resp:
        if resp.pricing_data:
            saved = resp.pricing_data
        if resp.raw_responses:
            site_profile = resp.raw_responses.get('site_profile', [])

    return jsonify({
        'success':      True,
        'appliances':   appliances,
        'sdwan':        sdwan,
        'segmentation': segmentation,
        'saved':        saved,          # { phase, rows: [...] } — manually saved pricing rows
        'site_profile': site_profile,   # site profile rows from assessment step
    })


# ── SAVE pricing inputs for this customer ────────────────────────────────────
@pricing_user_bp.route('/user/<customer_slug>/pricing-save', methods=['PATCH'])
def pricing_save(customer_slug):
    data = request.get_json() or {}
    resp = _slug_to_response(customer_slug)
    if not resp:
        return jsonify({'success': False, 'error': 'Customer not found'}), 404

    resp.pricing_data = {
        'phase':        data.get('phase', 'standard'),
        'rows':         data.get('rows', []),
        'support_pct':  data.get('support_pct', 20),
        'margin_pct':   data.get('margin_pct', 20),
        'services_amt': data.get('services_amt', 0),
    }
    db.session.commit()
    return jsonify({'success': True})


# ── Admin API: Site Profile Guidance text ─────────────────────────────────────
from models import AdminConfig

@pricing_bp.route('/admin/pricing/api/guidance', methods=['GET'])
def get_guidance():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    row = AdminConfig.query.filter_by(key='site_profile_guidance').first()
    return jsonify({
        'success': True,
        'value': row.value if row else (
            "Define your customer's site profile below. Each row represents a site tier "
            "(e.g. Small Branch, Factory, Data Centre). The size you select drives automatic "
            "SKU selection in the Pricing and TCO tabs. Check SD-WAN Required if the site "
            "needs an SD-WAN license, Segmentation Required if micro-segmentation is needed, "
            "and HA Required to select the dual-appliance SKU."
        )
    })



# ── Admin API: PPTX template info ─────────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/template-info', methods=['GET'])
def template_info():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    if not os.path.exists(TEMPLATE_PATH):
        return jsonify({'success': True, 'exists': False})
    meta = {}
    if os.path.exists(TEMPLATE_META):
        with open(TEMPLATE_META, 'r') as f:
            meta = json.load(f)
    return jsonify({'success': True, 'exists': True,
                    'filename': meta.get('filename', 'executive_template.pptx'),
                    'uploaded_at': meta.get('uploaded_at', '')})

# ── Admin API: PPTX template upload ──────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/template-upload', methods=['POST'])
def template_upload():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    f = request.files.get('file')
    if not f or not f.filename.endswith('.pptx'):
        return jsonify({'success': False, 'error': 'Please upload a .pptx file'}), 400
    original_name = secure_filename(f.filename)
    f.save(TEMPLATE_PATH)
    meta = {'filename': original_name,
            'uploaded_at': datetime.now(timezone.utc).strftime('%b %d, %Y %H:%M UTC')}
    with open(TEMPLATE_META, 'w') as mf:
        json.dump(meta, mf)
    return jsonify({'success': True})

# ── Admin API: PPTX template delete ──────────────────────────────────────────
@pricing_bp.route('/admin/pricing/api/template-delete', methods=['DELETE'])
def template_delete():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401
    for p in [TEMPLATE_PATH, TEMPLATE_META]:
        if os.path.exists(p):
            os.remove(p)
    return jsonify({'success': True})

@pricing_bp.route('/admin/pricing/api/guidance', methods=['PUT'])
def save_guidance():
    redir = admin_required()
    if redir:
        return jsonify({'success': False, 'error': 'Unauthorized'}), 401

    data  = request.get_json() or {}
    value = (data.get('value') or '').strip()

    row = AdminConfig.query.filter_by(key='site_profile_guidance').first()
    if row:
        row.value = value
    else:
        row = AdminConfig(key='site_profile_guidance', value=value)
        db.session.add(row)
    db.session.commit()
    return jsonify({'success': True})


# ── User: Generate Executive Summary PPTX ────────────────────────────────────
@pricing_user_bp.route('/user/<customer_slug>/export-pptx', methods=['GET'])
def export_pptx(customer_slug):
    if not os.path.exists(TEMPLATE_PATH):
        return jsonify({'success': False, 'error': 'No presentation template uploaded yet. Ask your admin to upload one at Admin → Pricing.'}), 404

    try:
        from pptx import Presentation
        from pptx.util import Pt
        import copy
    except ImportError:
        return jsonify({'success': False, 'error': 'python-pptx not installed on server. Run: pip install python-pptx'}), 500

    resp = _slug_to_response(customer_slug)
    if not resp:
        return jsonify({'success': False, 'error': 'Customer not found'}), 404

    # ── Gather data ────────────────────────────────────────────────────────
    customer_name = resp.customer_name or ''
    se_name       = resp.se_name       or ''
    today         = datetime.now().strftime('%B %d, %Y')

    # Pricing
    pricing       = resp.pricing_data or {}
    p_phase       = (pricing.get('phase') or 'standard').capitalize()
    p_rows        = pricing.get('rows', [])
    support_pct   = pricing.get('support_pct', 20)
    margin_pct    = pricing.get('margin_pct', 20)
    services_amt  = pricing.get('services_amt', 0)

    skus = PricingSKU.query.filter_by(active=True).all()
    def sku_price(sku_list, sid, phase):
        if not sid: return 0
        s = next((x for x in sku_list if x.id == int(sid)), None)
        return getattr(s, phase.lower(), 0) or 0 if s else 0

    al_subtotal = 0
    for row in p_rows:
        qty = row.get('qty', 0) or 0
        ap  = sku_price(skus, row.get('appliance_id'), p_phase)
        sw  = sku_price(skus, row.get('sdwan_id'),     p_phase)
        sg  = sku_price(skus, row.get('seg_id'),       p_phase)
        al_subtotal += qty * (ap + sw + sg)

    support_amt      = round(al_subtotal * support_pct / 100)
    margin_amt       = round(al_subtotal * margin_pct  / 100)
    annual_recurring = al_subtotal + support_amt + margin_amt
    yr1_total        = annual_recurring + (services_amt or 0)

    def fmt(v):
        if not v: return '$0'
        return '$' + f'{int(v):,}'

    # TCO / ROI  (tco_routes.py saves flat keys into raw_responses)
    raw_resp    = resp.raw_responses or {}
    tco_rows    = raw_resp.get('tco_rows',       []) or []
    fte_count   = raw_resp.get('tco_fte_count',   0) or 0
    fte_cost    = raw_resp.get('tco_fte_cost',    0) or 0
    breach_cost = raw_resp.get('tco_breach_cost', 0) or 0
    acv         = raw_resp.get('tco_acv_override') or annual_recurring
    legacy_ann  = 0
    from models import TCOEntry
    for trow in tco_rows:
        qty = trow.get('qty', 0) or 0
        for cat in ['fw_ns','sdwan','mpls','fw_ew','iot_ot','nac','l3sw','pam']:
            entry_id = trow.get(cat + '_id')
            if entry_id:
                entry = TCOEntry.query.get(int(entry_id))
                if entry:
                    legacy_ann += (entry.annual_cost or 0) * qty
    legacy_ann += fte_count * fte_cost + breach_cost
    annual_sav  = legacy_ann - acv
    yr3_sav     = annual_sav * 3
    roi_pct     = round(annual_sav / legacy_ann * 100) if legacy_ann > 0 else 0

    # Value drivers
    raw   = resp.raw_responses or {}
    vd_rows = []
    for sec in (resp.section_outcomes_cache or []) if hasattr(resp, 'section_outcomes_cache') else []:
        pass  # not available server-side easily; use raw_responses tag
    # Pull from outcomes stored in raw_responses if present
    outcomes = raw.get('_outcomes', [])

    # ── Tag substitution ──────────────────────────────────────────────────
    tags = {
        '{{customer_name}}': customer_name,
        '{{se_name}}':       se_name,
        '{{date}}':          today,
        '{{phase}}':         p_phase,
        '{{al_subtotal}}':   fmt(al_subtotal),
        '{{support_pct}}':   str(support_pct) + '%',
        '{{support_amt}}':   fmt(support_amt),
        '{{margin_pct}}':    str(margin_pct)  + '%',
        '{{margin_amt}}':    fmt(margin_amt),
        '{{annual_recurring}}': fmt(annual_recurring),
        '{{services_amt}}':  fmt(services_amt),
        '{{yr1_total}}':     fmt(yr1_total),
        '{{legacy_annual}}': fmt(legacy_ann),
        '{{annual_savings}}':fmt(annual_sav),
        '{{yr3_savings}}':   fmt(yr3_sav),
        '{{roi_pct}}':       (str(roi_pct) + '%') if legacy_ann > 0 else '—',
        '{{zscaler_acv}}':   fmt(acv),
    }

    def replace_tags(text):
        for k, v in tags.items():
            text = text.replace(k, str(v))
        return text

    def process_shape(shape):
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    run.text = replace_tags(run.text)
        if shape.shape_type == 6:  # GROUP
            for s in shape.shapes:
                process_shape(s)

    prs = Presentation(TEMPLATE_PATH)
    for slide in prs.slides:
        for shape in slide.shapes:
            process_shape(shape)

    import io
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    safe_name = customer_name.replace(' ', '_').replace('/', '-')
    filename  = f'ZTB_Executive_Summary_{safe_name}_{datetime.now().strftime("%Y%m%d")}.pptx'

    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation')

