# admin_responses_routes.py — Admin Responses Module
from flask import Blueprint, request, jsonify, session, send_file
from models import db, UserResponse, Question
from functools import wraps
import json, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

responses_bp = Blueprint('responses', __name__, url_prefix='/admin/responses')

def require_admin(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get('admin_logged_in'):
            return jsonify({'success': False, 'error': 'Unauthorized'}), 401
        return f(*args, **kwargs)
    return decorated

def _get_keyword(question):
    """keyword = category name, matching how admin_questions_routes stores it."""
    if question.category_ref and question.category_ref.name:
        name = question.category_ref.name.strip()
        if name and name.lower() != 'general':
            return name
    return f'q_{question.id}'

# ── GET /admin/responses/list ─────────────────────────────────────────────
@responses_bp.route('/list', methods=['GET'])
@require_admin
def list_responses():
    search   = request.args.get('search', '').strip().lower()
    customer = request.args.get('customer', '').strip().lower()
    se       = request.args.get('se', '').strip().lower()
    kw_key   = request.args.get('keyword', '').strip().lower()
    kw_val   = request.args.get('value', '').strip().lower()

    all_rows = UserResponse.query.order_by(UserResponse.completed_at.desc()).all()
    results = []

    for r in all_rows:
        answers = r.answers or {}

        if customer and customer not in (r.customer_name or '').lower():
            continue
        if se and se not in (r.se_name or '').lower():
            continue
        if kw_key:
            cell = answers.get(kw_key)
            if cell is None:
                continue
            if kw_val:
                cell_str = ', '.join(cell) if isinstance(cell, list) else str(cell)
                if kw_val not in cell_str.lower():
                    continue
        if search:
            blob = json.dumps(answers).lower() + (r.customer_name or '').lower() + (r.se_name or '').lower()
            if search not in blob:
                continue

        results.append({
            'id':            r.id,
            'customer_name': r.customer_name,
            'se_name':       r.se_name,
            'completed_at':  r.completed_at.strftime('%Y-%m-%d %H:%M') if r.completed_at else '',
            'answers':       answers
        })

    return jsonify({'success': True, 'responses': results, 'total': len(results)})

# ── GET /admin/responses/keywords ─────────────────────────────────────────
@responses_bp.route('/keywords', methods=['GET'])
@require_admin
def list_keywords():
    questions = Question.query.order_by(Question.order).all()
    keywords = []
    seen = set()
    for q in questions:
        if q.info_only:
            continue
        kw = _get_keyword(q)
        if kw not in seen:
            seen.add(kw)
            keywords.append({'keyword': kw, 'text': q.text})
    return jsonify({'success': True, 'keywords': keywords})

# ── DELETE /admin/responses/<id> ──────────────────────────────────────────
@responses_bp.route('/<int:response_id>', methods=['DELETE'])
@require_admin
def delete_response(response_id):
    r = UserResponse.query.get_or_404(response_id)
    db.session.delete(r)
    db.session.commit()
    return jsonify({'success': True, 'message': f'Response {response_id} deleted'})

# ── POST /admin/responses/bulk-delete ────────────────────────────────────
@responses_bp.route('/bulk-delete', methods=['POST'])
@require_admin
def bulk_delete():
    data = request.get_json()
    ids = data.get('ids', [])
    if not ids:
        return jsonify({'success': False, 'error': 'No IDs provided'}), 400
    deleted = 0
    for rid in ids:
        r = UserResponse.query.get(rid)
        if r:
            db.session.delete(r)
            deleted += 1
    db.session.commit()
    return jsonify({'success': True, 'deleted': deleted})

# ── GET /admin/responses/export ───────────────────────────────────────────
@responses_bp.route('/export', methods=['GET'])
@require_admin
def export_responses():
    search   = request.args.get('search', '').strip().lower()
    customer = request.args.get('customer', '').strip().lower()
    se       = request.args.get('se', '').strip().lower()
    kw_key   = request.args.get('keyword', '').strip().lower()
    kw_val   = request.args.get('value', '').strip().lower()

    all_rows = UserResponse.query.order_by(UserResponse.completed_at.desc()).all()
    filtered = []

    for r in all_rows:
        answers = r.answers or {}
        if customer and customer not in (r.customer_name or '').lower():
            continue
        if se and se not in (r.se_name or '').lower():
            continue
        if kw_key:
            cell = answers.get(kw_key)
            if cell is None:
                continue
            if kw_val:
                cell_str = ', '.join(cell) if isinstance(cell, list) else str(cell)
                if kw_val not in cell_str.lower():
                    continue
        if search:
            blob = json.dumps(answers).lower() + (r.customer_name or '').lower() + (r.se_name or '').lower()
            if search not in blob:
                continue
        filtered.append(r)

    # Build keyword list from questions in order
    questions = Question.query.order_by(Question.order).all()
    all_keywords = []
    seen_kw = set()
    for q in questions:
        if q.info_only:
            continue
        kw = _get_keyword(q)
        if kw not in seen_kw:
            all_keywords.append((kw, q.text))
            seen_kw.add(kw)

    # Build Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Responses"

    hdr_font  = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill  = PatternFill("solid", fgColor="003366")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    side      = Side(style="thin", color="CCCCCC")
    border    = Border(left=side, right=side, top=side, bottom=side)

    fixed_cols = ['ID', 'Customer Name', 'SE Name', 'Submitted At']
    all_cols   = fixed_cols + [q_text for (kw, q_text) in all_keywords]

    for col_idx, label in enumerate(all_cols, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 30

    for row_idx, r in enumerate(filtered, start=2):
        answers = r.answers or {}
        row_data = [
            r.id,
            r.customer_name or '',
            r.se_name or '',
            r.completed_at.strftime('%Y-%m-%d %H:%M') if r.completed_at else ''
        ]
        for (kw, q_text) in all_keywords:
            val = answers.get(kw, '')
            if isinstance(val, list):
                val = ', '.join(val)
            row_data.append(val)

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border    = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F6FC")
        ws.row_dimensions[row_idx].height = 20

    for col_cells in ws.columns:
        length = max((len(str(c.value)) if c.value else 0) for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 12), 50)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"ZTB_Responses_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )
