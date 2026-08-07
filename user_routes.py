# user_routes.py — Complete User-Facing Routes for ZTB Super App
from flask import Blueprint, request, jsonify, session, render_template, send_file, redirect, url_for
from models import (db, UserResponse, Question, QuestionOption,
                    AssessmentConfig, QuestionCategory,
                    ValueProp, Asset, TestCase, POVPlanner, Roadblock,
                    ColumnDefinition)
import io
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import text as sa_text
from datetime import datetime

user_bp = Blueprint('user', __name__, url_prefix='/user')

# ── Formatting helpers ────────────────────────────────────────────────────────

def _hdr(ws, row, col, value, bg="003366", fg="FFFFFF", bold=True, size=11):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(color=fg, bold=bold, size=size)
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return cell

def _thin_border():
    side = Side(style="thin", color="CCCCCC")
    return Border(left=side, right=side, top=side, bottom=side)

def _autosize(ws, min_w=10, max_w=60):
    for col_cells in ws.columns:
        length = max(
            (len(str(c.value)) if c.value is not None else 0)
            for c in col_cells
        )
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(
            max(length + 2, min_w), max_w
        )

def _get_col_defs(table_name):
    defs = (ColumnDefinition.query
            .filter_by(table_name=table_name)
            .order_by(ColumnDefinition.display_order)
            .all())
    return [(d.column_key, d.column_label) for d in defs]

def _collect(table_name, session_ids=None):
    results = []
    seen_ids = set()
    try:
        with db.engine.connect() as conn:
            rows = conn.execute(sa_text(
                f"SELECT * FROM {table_name} WHERE universal = 'yes'"
            )).fetchall()
            if rows:
                keys = rows[0]._fields
                for r in rows:
                    d = dict(zip(keys, r))
                    if d['id'] not in seen_ids:
                        seen_ids.add(d['id'])
                        results.append(d)
            if session_ids:
                placeholders = ','.join(str(int(i)) for i in session_ids)
                rows2 = conn.execute(sa_text(
                    f"SELECT * FROM {table_name} WHERE id IN ({placeholders})"
                )).fetchall()
                if rows2:
                    keys = rows2[0]._fields
                    for r in rows2:
                        d = dict(zip(keys, r))
                        if d['id'] not in seen_ids:
                            seen_ids.add(d['id'])
                            results.append(d)
    except Exception as e:
        print(f"_collect error for {table_name}: {e}")
    return results

def _write_tab(ws, rows, col_defs):
    if not col_defs:
        ws.cell(row=1, column=1, value="No columns defined.")
        return
    for ci, (key, label) in enumerate(col_defs, start=1):
        _hdr(ws, 1, ci, label)
    ws.row_dimensions[1].height = 30
    border = _thin_border()
    for ri, row_dict in enumerate(rows, start=2):
        for ci, (key, label) in enumerate(col_defs, start=1):
            val = row_dict.get(key, '')
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if ri % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F0F4FF")
    ws.freeze_panes = "A2"
    _autosize(ws)

def _get_keyword(question):
    if question.category_ref and question.category_ref.name:
        name = question.category_ref.name.strip()
        if name and name.lower() != 'general':
            return name
    return f'q_{question.id}'

def _make_slug(name):
    """'Acme Corp' -> 'acme-corp'"""
    slug = (name or 'unknown').strip().lower()
    slug = re.sub(r'[^a-z0-9\s-]', '', slug)
    slug = re.sub(r'\s+', '-', slug)
    slug = re.sub(r'-+', '-', slug).strip('-')
    return slug or 'unknown'

def _run_mapping(responses):
    """Re-usable mapping logic: responses dict -> results payload + keyword_answers."""
    vp_ids    = set()
    asset_ids = set()
    tc_ids    = set()
    rb_ids    = set()
    pov_ids   = set()
    keyword_answers = {}

    for q_id_str, value in responses.items():
        try:
            q_id = int(q_id_str)
        except (ValueError, TypeError):
            continue
        question = Question.query.get(q_id)
        if not question or question.info_only:
            continue

        kw = _get_keyword(question)
        selected_option_ids = []

        if question.options_type in ('select_all', 'Select All'):
            if isinstance(value, list):
                selected_option_ids = [int(v) for v in value if str(v).isdigit()]
            labels = []
            for opt_id in selected_option_ids:
                opt = QuestionOption.query.get(opt_id)
                if opt:
                    labels.append(opt.label)
            keyword_answers[kw] = labels

        elif question.options_type in ('text', 'Text', 'Textbox', 'textbox'):
            keyword_answers[kw] = value

        else:
            if value and str(value).isdigit():
                selected_option_ids = [int(value)]
            opt = QuestionOption.query.get(selected_option_ids[0]) if selected_option_ids else None
            keyword_answers[kw] = opt.label if opt else (value or '')

        for opt_id in selected_option_ids:
            opt = QuestionOption.query.get(opt_id)
            if not opt:
                continue
            for a in opt.assets:       asset_ids.add(a.id)
            for v in opt.value_props:  vp_ids.add(v.id)
            for t in opt.test_cases:   tc_ids.add(t.id)
            for p in opt.pov_steps:    pov_ids.add(p.id)
            for r in opt.roadblocks:   rb_ids.add(r.id)

    results_payload = {
        'vp_ids':    list(vp_ids),
        'asset_ids': list(asset_ids),
        'tc_ids':    list(tc_ids),
        'pov_ids':   list(pov_ids),
        'rb_ids':    list(rb_ids),
    }
    return results_payload, keyword_answers


# ── Fixed routes — all defined before /<customer_slug> ───────────────────────

@user_bp.route('/')
def landing():
    in_progress = (UserResponse.query
                   .filter_by(status='in_progress')
                   .order_by(UserResponse.started_at.desc())
                   .limit(15)
                   .all())
    total_questions = Question.query.filter_by(info_only=False).count()
    return render_template('user_landing.html',
                           in_progress=in_progress,
                           total_questions=total_questions)

@user_bp.route('/start', methods=['POST'])
def start_assessment():
    if request.is_json:
        data = request.get_json()
    else:
        data = request.form.to_dict()
    customer_name   = (data.get('customer_name') or '').strip()
    se_name         = (data.get('se_name') or '').strip()
    opportunity_url = (data.get('opportunity_url') or '').strip()
    if not customer_name:
        return jsonify({'error': 'Customer name is required'}), 400
    session['customer_name']      = customer_name
    session['se_name']            = se_name
    session['opportunity_url']    = opportunity_url
    session['assessment_started'] = True
    session['responses']          = {}
    session['draft_response_id']  = None
    return jsonify({'success': True, 'message': 'Assessment started'})

@user_bp.route('/questions', methods=['GET'])
def get_questions():
    all_questions = Question.query.order_by(Question.order).all()
    result = []
    for q in all_questions:
        cat = QuestionCategory.query.get(q.category_id)
        opts = []
        for o in sorted(q.options, key=lambda x: x.id):
            opts.append({'id': o.id, 'label': o.label})
        result.append({
            'question_id':        q.id,
            'category':           cat.name if cat else '',
            'category_type_name': q.category_type_ref.name if q.category_type_ref else '',
            'text':               q.text,
            'options_type':       q.options_type,
            'info_only':          q.info_only,
            'options':            opts
        })
    return render_template('user_assessment.html',
                           questions=result,
                           saved_responses={},
                           resume_index=0,
                           edit_mode=False,
                           customer_slug='')

@user_bp.route('/save', methods=['POST'])
def save_progress():
    data          = request.get_json()
    raw_responses = data.get('responses', {})
    current_index = data.get('current_index', 0)

    customer_name   = session.get('customer_name') or data.get('customer_name', 'Unknown')
    se_name         = session.get('se_name')        or data.get('se_name', '')
    opportunity_url = session.get('opportunity_url', '')

    draft_id = session.get('draft_response_id')
    if draft_id:
        user_resp = UserResponse.query.get(draft_id)
    else:
        user_resp = None

    if user_resp and user_resp.status == 'in_progress':
        user_resp.raw_responses          = raw_responses
        user_resp.current_question_index = current_index
    else:
        user_resp = UserResponse(
            customer_name          = customer_name,
            se_name                = se_name,
            opportunity_url        = opportunity_url,
            status                 = 'in_progress',
            raw_responses          = raw_responses,
            current_question_index = current_index,
            answers                = {},
            results                = {}
        )
        db.session.add(user_resp)

    db.session.commit()
    session['draft_response_id'] = user_resp.id

    return jsonify({'success': True, 'response_id': user_resp.id})

@user_bp.route('/resume/<int:response_id>', methods=['GET'])
def resume_assessment(response_id):
    user_resp = UserResponse.query.get_or_404(response_id)

    if user_resp.status != 'in_progress':
        return redirect(url_for('user.landing'))

    session['customer_name']      = user_resp.customer_name
    session['se_name']            = user_resp.se_name
    session['opportunity_url']    = user_resp.opportunity_url
    session['assessment_started'] = True
    session['responses']          = {}
    session['draft_response_id']  = user_resp.id

    all_questions = Question.query.order_by(Question.order).all()
    result = []
    for q in all_questions:
        cat = QuestionCategory.query.get(q.category_id)
        opts = []
        for o in sorted(q.options, key=lambda x: x.id):
            opts.append({'id': o.id, 'label': o.label})
        result.append({
            'question_id':        q.id,
            'category':           cat.name if cat else '',
            'category_type_name': q.category_type_ref.name if q.category_type_ref else '',
            'text':               q.text,
            'options_type':       q.options_type,
            'info_only':          q.info_only,
            'options':            opts
        })

    return render_template('user_assessment.html',
                           questions=result,
                           saved_responses=user_resp.raw_responses or {},
                           resume_index=user_resp.current_question_index or 0,
                           edit_mode=False,
                           customer_slug='')

@user_bp.route('/submit', methods=['POST'])
def submit_responses():
    data      = request.get_json()
    responses = data.get('responses', {})
    customer_name   = session.get('customer_name') or data.get('customer_name', 'Unknown')
    se_name         = session.get('se_name')        or data.get('se_name', '')
    opportunity_url = session.get('opportunity_url', '')

    results_payload, keyword_answers = _run_mapping(responses)

    # Keep session populated for legacy /user/export
    session['vp_ids']    = results_payload['vp_ids']
    session['asset_ids'] = results_payload['asset_ids']
    session['tc_ids']    = results_payload['tc_ids']
    session['rb_ids']    = results_payload['rb_ids']
    session['pov_ids']   = results_payload['pov_ids']

    draft_id  = session.get('draft_response_id')
    user_resp = UserResponse.query.get(draft_id) if draft_id else None

    if user_resp and user_resp.status == 'in_progress':
        user_resp.answers       = keyword_answers
        user_resp.raw_responses = responses
        user_resp.results       = results_payload
        user_resp.status        = 'completed'
        user_resp.completed_at  = datetime.utcnow()
    else:
        user_resp = UserResponse(
            customer_name   = customer_name,
            se_name         = se_name,
            opportunity_url = opportunity_url,
            answers         = keyword_answers,
            raw_responses   = responses,
            results         = results_payload,
            status          = 'completed',
            completed_at    = datetime.utcnow()
        )
        db.session.add(user_resp)

    db.session.commit()
    session['response_id']       = user_resp.id
    session['draft_response_id'] = None

    slug = _make_slug(customer_name)
    return jsonify({'success': True, 'response_id': user_resp.id, 'customer_slug': slug})

@user_bp.route('/export', methods=['GET'])
def export_excel():
    response_id     = session.get('response_id')
    customer_name   = session.get('customer_name', 'Unknown')
    se_name         = session.get('se_name', '')
    opportunity_url = session.get('opportunity_url', '')

    user_resp = UserResponse.query.get(response_id) if response_id else None
    answers   = user_resp.answers if user_resp else {}
    if user_resp and user_resp.opportunity_url:
        opportunity_url = user_resp.opportunity_url

    wb = Workbook()

    ws_assess = wb.active
    ws_assess.title = "Assessment"
    ws_assess.merge_cells("A1:B1")
    cell = ws_assess["A1"]
    cell.value = f"Customer: {customer_name}"
    cell.font = Font(bold=True, size=14, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="003366")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws_assess.row_dimensions[1].height = 30
    ws_assess.merge_cells("A2:B2")
    cell2 = ws_assess["A2"]
    cell2.value = f"SE: {se_name}"
    cell2.font = Font(bold=True, size=11, color="FFFFFF")
    cell2.fill = PatternFill("solid", fgColor="003366")
    cell2.alignment = Alignment(horizontal="left", vertical="center")
    ws_assess.row_dimensions[2].height = 22
    ws_assess.merge_cells("A3:B3")
    cell3 = ws_assess["A3"]
    cell3.value = f"Opportunity URL: {opportunity_url}" if opportunity_url else "Opportunity URL: —"
    cell3.font = Font(bold=False, size=10, color="FFFFFF")
    cell3.fill = PatternFill("solid", fgColor="003366")
    cell3.alignment = Alignment(horizontal="left", vertical="center")
    ws_assess.row_dimensions[3].height = 20
    ws_assess.row_dimensions[4].height = 10
    _hdr(ws_assess, 5, 1, "Question")
    _hdr(ws_assess, 5, 2, "Answer")
    ws_assess.row_dimensions[5].height = 28
    ws_assess.column_dimensions["A"].width = 55
    ws_assess.column_dimensions["B"].width = 45
    border = _thin_border()
    current_row = 6
    all_questions = Question.query.order_by(Question.order).all()
    for q in all_questions:
        if q.info_only:
            continue
        kw = _get_keyword(q)
        raw_val = answers.get(kw, '')
        if isinstance(raw_val, list):
            answer_text = ', '.join(raw_val)
        else:
            answer_text = str(raw_val) if raw_val else '—'
        q_cell = ws_assess.cell(row=current_row, column=1, value=q.text)
        a_cell = ws_assess.cell(row=current_row, column=2, value=answer_text)
        for c in (q_cell, a_cell):
            c.border = border
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if current_row % 2 == 0:
            for c in (q_cell, a_cell):
                c.fill = PatternFill("solid", fgColor="F0F4FF")
        current_row += 1
    ws_assess.freeze_panes = "A6"

    ws_vp = wb.create_sheet("Value Props")
    _write_tab(ws_vp, _collect("value_props", session.get('vp_ids', [])), _get_col_defs("value_props"))
    ws_assets = wb.create_sheet("Assets")
    _write_tab(ws_assets, _collect("assets", session.get('asset_ids', [])), _get_col_defs("assets"))
    ws_tc = wb.create_sheet("Test Cases")
    _write_tab(ws_tc, _collect("test_cases", session.get('tc_ids', [])), _get_col_defs("test_cases"))
    ws_pov = wb.create_sheet("POV Planner")
    _write_tab(ws_pov, _collect("pov_planner", session.get('pov_ids', [])), _get_col_defs("pov_planner"))
    ws_rb = wb.create_sheet("Roadblocks")
    _write_tab(ws_rb, _collect("roadblocks", session.get('rb_ids', [])), _get_col_defs("roadblocks"))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"ZTB_Assessment_{customer_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=filename)

@user_bp.route('/bom')
def bom():
    return render_template('user_bom.html')


# ── Variable routes — MUST stay last in this file ────────────────────────────

@user_bp.route('/<customer_slug>')
def customer_results(customer_slug):
    """Permanent results page. Always shows most recent completed assessment for this slug."""
    completed = (UserResponse.query
                 .filter_by(status='completed')
                 .order_by(UserResponse.completed_at.desc())
                 .all())

    user_resp = None
    for r in completed:
        if _make_slug(r.customer_name) == customer_slug:
            user_resp = r
            break

    if not user_resp:
        return render_template('user_results.html',
                               not_found=True,
                               customer_slug=customer_slug), 404

    results   = user_resp.results or {}
    vp_ids    = results.get('vp_ids', [])
    asset_ids = results.get('asset_ids', [])

    value_props = _collect("value_props", vp_ids)
    assets      = _collect("assets", asset_ids)

    # Build ordered Q&A pairs
    all_questions = Question.query.order_by(Question.order).all()
    answers  = user_resp.answers or {}
    qa_pairs = []
    for q in all_questions:
        if q.info_only:
            continue
        kw    = _get_keyword(q)
        value = answers.get(kw)
        if value is None:
            continue
        if isinstance(value, list):
            display = ', '.join(value) if value else '—'
        else:
            display = str(value).strip() if value else '—'
        cat = (q.category_type_ref.name if q.category_type_ref else
               q.category_ref.name if q.category_ref else '')
        qa_pairs.append({
            'category': cat,
            'question': q.text,
            'answer':   display,
        })

    return render_template('user_results.html',
                           not_found=False,
                           user_resp=user_resp,
                           customer_slug=customer_slug,
                           value_props=value_props,
                           assets=assets,
                           qa_pairs=qa_pairs)


@user_bp.route('/<customer_slug>/edit')
def edit_assessment(customer_slug):
    """Load a completed assessment pre-filled for editing."""
    completed = (UserResponse.query
                 .filter_by(status='completed')
                 .order_by(UserResponse.completed_at.desc())
                 .all())

    user_resp = None
    for r in completed:
        if _make_slug(r.customer_name) == customer_slug:
            user_resp = r
            break

    if not user_resp:
        return redirect(url_for('user.landing'))

    # Populate session so /update knows who this is
    session['customer_name']      = user_resp.customer_name
    session['se_name']            = user_resp.se_name
    session['opportunity_url']    = user_resp.opportunity_url
    session['edit_response_id']   = user_resp.id

    all_questions = Question.query.order_by(Question.order).all()
    result = []
    for q in all_questions:
        cat = QuestionCategory.query.get(q.category_id)
        opts = []
        for o in sorted(q.options, key=lambda x: x.id):
            opts.append({'id': o.id, 'label': o.label})
        result.append({
            'question_id':        q.id,
            'category':           cat.name if cat else '',
            'category_type_name': q.category_type_ref.name if q.category_type_ref else '',
            'text':               q.text,
            'options_type':       q.options_type,
            'info_only':          q.info_only,
            'options':            opts
        })

    return render_template('user_assessment.html',
                           questions=result,
                           saved_responses=user_resp.raw_responses or {},
                           resume_index=0,
                           edit_mode=True,
                           customer_slug=customer_slug)


@user_bp.route('/<customer_slug>/update', methods=['POST'])
def update_assessment(customer_slug):
    """Re-run mapping on updated responses and save back to the same record."""
    data      = request.get_json()
    responses = data.get('responses', {})

    edit_id   = session.get('edit_response_id')
    user_resp = UserResponse.query.get(edit_id) if edit_id else None

    # Fallback: find by slug if session was lost
    if not user_resp:
        completed = (UserResponse.query
                     .filter_by(status='completed')
                     .order_by(UserResponse.completed_at.desc())
                     .all())
        for r in completed:
            if _make_slug(r.customer_name) == customer_slug:
                user_resp = r
                break

    if not user_resp:
        return jsonify({'error': 'Assessment not found'}), 404

    results_payload, keyword_answers = _run_mapping(responses)

    user_resp.raw_responses = responses
    user_resp.answers       = keyword_answers
    user_resp.results       = results_payload
    user_resp.completed_at  = datetime.utcnow()
    db.session.commit()

    session['edit_response_id'] = None

    return jsonify({'success': True, 'customer_slug': customer_slug})


@user_bp.route('/<customer_slug>/export')
def customer_export(customer_slug):
    """POV Planner + Test Cases Excel scoped to a specific customer."""
    completed = (UserResponse.query
                 .filter_by(status='completed')
                 .order_by(UserResponse.completed_at.desc())
                 .all())

    user_resp = None
    for r in completed:
        if _make_slug(r.customer_name) == customer_slug:
            user_resp = r
            break

    if not user_resp:
        return jsonify({'error': 'Not found'}), 404

    results = user_resp.results or {}
    tc_ids  = results.get('tc_ids', [])
    pov_ids = results.get('pov_ids', [])

    wb = Workbook()

    ws_pov = wb.active
    ws_pov.title = "POV Planner"
    _write_tab(ws_pov, _collect("pov_planner", pov_ids), _get_col_defs("pov_planner"))

    ws_tc = wb.create_sheet("Test Cases")
    _write_tab(ws_tc, _collect("test_cases", tc_ids), _get_col_defs("test_cases"))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"ZTB_POV_{user_resp.customer_name.replace(' ', '_')}_{datetime.utcnow().strftime('%Y%m%d')}.xlsx"
    return send_file(output,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=filename)
